"""Real-text extraction from ACT's PAG written-part DOCX documents (see
DECISIONS.md / PAG_PAP_SPEC.md). Verified live against real documents: they
open with a real legal article heading ("Art. 5 Zone mixte urbaine
[MIX-u]", "Art.19 Zone forestière [FOR]" — note the inconsistent spacing
after "Art." in the real documents, handled below), followed by paragraphs
and, sometimes, a table of exceptions/overrides.
"""

from __future__ import annotations

import io
import re

from docx import Document as DocxDocument

_ARTICLE_RE = re.compile(r"^Art\.\s*(\d+)\b")


def extract_pag_docx(data: bytes) -> tuple[str, str, str | None]:
    """Returns (title, full_text, article_ref) — full_text includes any
    table content (verbatim, cell-joined) since real coefficient/override
    tables live there (see the MIX_u schéma-directeur table in
    DECISIONS.md), not just prose paragraphs."""
    doc = DocxDocument(io.BytesIO(data))
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    title = paragraphs[0] if paragraphs else ""

    parts = list(paragraphs)
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    full_text = "\n".join(parts)

    match = _ARTICLE_RE.match(title)
    article_ref = f"Art. {match.group(1)}" if match else None
    return title, full_text, article_ref
